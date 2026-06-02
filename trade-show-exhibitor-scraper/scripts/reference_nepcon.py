#!/usr/bin/env python3
"""
NEPCON Japan 2026 Exhibitor Scraper - Final v4
Uses GraphQL API with browser-extracted cookies for auth.
"""
import json, time, asyncio, httpx
from pathlib import Path

EVENT_EDITION_ID = "eve-c9ed2b9d-84a3-4932-ab81-e3e1d1441041"
GRAPHQL_URL = "https://api.reedexpo.com/graphql/"
LINKS_JSON = Path(__file__).parent / "links_raw.json"
OUTPUT_XLSX = Path(__file__).parent / "nepcon_japan_2026_exhibitors.xlsx"
CONCURRENCY = 30

AUTH_COOKIES = {
    "ClientId": "uhQVcmxLwXAjVtVpTvoerERiZSsNz0om",
    "id_token": "eyJhbGciOiJSUzI1NiIsImtpZCI6IjIxQTZDNDJGOTE3MzE3QzBBODdGODBCQzVFMDQ0NjA0MEUwNjlBRDdSUzI1NiIsIng1dCI6IklhYkVMNUZ6RjhDb2Y0QzhYZ1JHQkE0R210YyIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJodHRwczovL2F1dGgucmVlZGV4cG8uY29tL3NlY3VyZSIsIm5iZiI6MTc4MDI3OTg4MSwiaWF0IjoxNzgwMjc5ODgxLCJleHAiOjE3OTUyNzk4ODEsImF1ZCI6WyJ1cm46cng6ZGlnaXRhbDphcGk6d2F0Y2hib3giLCJodHRwczovL2F1dGgucmVlZGV4cG8uY29tL3NlY3VyZS9yZXNvdXJjZXMiXSwic2NvcGUiOlsidXJuOnJ4OmRpZ2l0YWw6YWN0aW9uOndyaXRlIl0sImFtciI6WyJwYXNzd29yZCJdLCJjbGllbnRfaWQiOiI1ZmRmNjRiODQyNzc0ODM4OTc2YTUzZjcwYWI2MWNjNyIsInN1YiI6IjUzNGRlYWM3MjNiODQxZTg4ZTRlMmJmZjU4NzUwYzQ1IiwiYXV0aF90aW1lIjoxNzgwMjc5ODgxLCJpZHAiOiJsb2NhbCIsInJvbGUiOiJhbm9ueW1vdXMifQ.oMWGq3NOaWCBXEJLn1dbZzB2pL47SAoYr8DSrDGdKm3gRUAI4a6d8kO_u8qCUwBst0OnUnV1hMyfzs2mv1xwMbfdhzw0jt_v7bWjbQqS1uwdAdplaTpGGeqIuR2sJ5P_5btW51wJO5axguPbw6JjF8NO_tyKVuaWjagsHTlRk6Fzr8-OZ1XbYQRddJmm3OIsBat2jXDcm2KEp3c5vZZ1J3gZJhkJF9vo7Ofnvf4jcQyUvLCSxM3SThf-zDIe6fDaeUS8jVwZFnH_hN2EkhXg1I8oaf81gd5OazNo-W1jwgWseyRFVw-6FsrDhpXnJOcVaW2OkYnfjZRdNzWBbHbogw",
}

COMMON_HEADERS = {
    "x-clientid": "uhQVcmxLwXAjVtVpTvoerERiZSsNz0om",
    "Content-Type": "application/json",
    "Accept": "application/json",
    "Origin": "https://www.nepconjapan.jp",
    "Referer": "https://www.nepconjapan.jp/",
}


def load_links():
    with open(LINKS_JSON, "r", encoding="utf-8-sig") as f:
        raw = json.load(f)
    if isinstance(raw, str): raw = json.loads(raw)
    items = [{"name": d["n"], "orgGuid": d["g"], "url": d["u"]} for d in raw]
    print(f"Loaded {len(items)} exhibitors")
    return items


async def fetch_one(client, item, sem):
    org = item["orgGuid"]
    q = '{ exhibitingOrganisation( eventEditionId: "%s" organisationId: "%s" ){ id organisationId isNew companyName contactEmail website phone socialMedia{url name} multilingual{logoUrl displayName description showObjective representedBrands{name} addressLine1 stateProvince city country countryCode postcode} filterCategories{multilingual{name locale}responses{multilingual{name locale}}} products(includeUnpublished:false){id imageUrl multilingual{title description locale}} stands{name} } }' % (
        EVENT_EDITION_ID, "org-" + org
    )
    
    for _ in range(3):
        try:
            async with sem:
                r = await client.post(GRAPHQL_URL, json={"query": q}, timeout=25)
            if r.status_code == 200:
                d = r.json().get("data", {}).get("exhibitingOrganisation")
                if d: return parse(d, item["name"])
            await asyncio.sleep(1)
        except: await asyncio.sleep(0.5)
    return None


def parse(d, name):
    en = None
    for m in d.get("multilingual", []):
        if m.get("locale", "").lower() == "en-gb": en = m; break
    if not en and d.get("multilingual"): en = d["multilingual"][0]
    
    cats, prods, brands, soc = [], [], [], {}
    for c in d.get("filterCategories", []):
        for r in c.get("responses", []):
            for ml in r.get("multilingual", []):
                if ml.get("locale") == "en-GB": cats.append(ml["name"])
    for p in d.get("products", []):
        for ml in p.get("multilingual", []): prods.append(ml.get("title", ""))
    for sm in d.get("socialMedia", []): soc[sm["name"].upper()] = sm["url"]
    if en:
        for b in en.get("representedBrands", []): brands.append(b.get("name", ""))
    stands = [s["name"] for s in d.get("stands", [])]
    
    return {
        "company_name": d.get("companyName", ""),
        "display_name": en.get("displayName", "") if en else "",
        "stand_number": ", ".join(stands),
        "website": d.get("website", ""), "email": d.get("contactEmail", ""),
        "phone": d.get("phone", ""),
        "country": en.get("country", "") if en else "",
        "country_code": en.get("countryCode", "") if en else "",
        "city": en.get("city", "") if en else "",
        "state_province": en.get("stateProvince", "") if en else "",
        "address": en.get("addressLine1", "") if en else "",
        "postcode": en.get("postcode", "") if en else "",
        "description": en.get("description", "") if en else "",
        "show_objective": en.get("showObjective", "") if en else "",
        "categories": ", ".join(cats), "brands": ", ".join(brands),
        "products": ", ".join(prods), "products_count": len(d.get("products", [])),
        "facebook": soc.get("FACEBOOK", ""), "twitter": soc.get("TWITTER", ""),
        "linkedin": soc.get("LINKEDIN", ""), "youtube": soc.get("YOUTUBE", ""),
        "instagram": soc.get("INSTAGRAM", ""),
        "is_new": "Yes" if d.get("isNew") else "No",
        "org_id": d.get("organisationId", ""),
    }


def sanitize(val):
    """Remove illegal characters for Excel"""
    import re
    if not isinstance(val, str): return val
    # Remove control characters except newline, carriage return, tab
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', val)


def save_xlsx(results, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    
    wb = Workbook(); ws = wb.active; ws.title = "NEPCON Japan 2026"
    headers = ["Company Name","Display Name","Stand Number","Country","Country Code","City","State/Province","Address","Postcode","Website","Email","Phone","Categories","Brands","Products Count","Products","Description","Show Objective","Facebook","Twitter","LinkedIn","YouTube","Instagram","Is New","Org ID"]
    fields = ["company_name","display_name","stand_number","country","country_code","city","state_province","address","postcode","website","email","phone","categories","brands","products_count","products","description","show_objective","facebook","twitter","linkedin","youtube","instagram","is_new","org_id"]
    
    hf = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill(start_color="005BAC", end_color="005BAC", fill_type="solid")
    ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
    tb = Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
    
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h); cell.font = hf; cell.fill = hfill; cell.alignment = ha; cell.border = tb
    for ri, item in enumerate(results, 2):
        for ci, f in enumerate(fields, 1):
            val = sanitize(item.get(f, ""))
            cell = ws.cell(row=ri, column=ci, value=val); cell.font = Font(size=10); cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = tb
    
    widths = [35,35,15,20,12,20,20,40,12,30,25,20,30,25,12,50,50,50,30,30,30,30,30,8,40]
    for i, w in enumerate(widths): ws.column_dimensions[chr(65+i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "NEPCON Japan 2026 Exhibitors"; ws2["A1"].font = Font(bold=True, size=14)
    ws2["A3"], ws2["B3"] = "Total:", len(results)
    ws2["A4"], ws2["B4"] = "Generated:", time.strftime("%Y-%m-%d %H:%M:%S")
    cts = {}
    for r in results: c = r.get("country","Unknown") or "Unknown"; cts[c] = cts.get(c,0)+1
    ws2["A6"] = "Country Breakdown"; ws2["A6"].font = Font(bold=True, size=12)
    ws2["A7"], ws2["B7"] = "Country", "Count"
    for i, (c, n) in enumerate(sorted(cts.items(), key=lambda x:-x[1]), 8): ws2[f"A{i}"], ws2[f"B{i}"] = c, n
    wb.save(path)
    print(f"\nSaved {len(results)} exhibitors to {path}")


async def main():
    items = load_links()
    if not items: print("No links!"); return
    
    sem = asyncio.Semaphore(CONCURRENCY)
    results = []; total = len(items)
    
    async with httpx.AsyncClient(headers=COMMON_HEADERS, cookies=AUTH_COOKIES, follow_redirects=True) as cli:
        print(f"\nFetching {total} exhibitors ({CONCURRENCY} concurrent)...")
        t0 = time.time()
        tasks = [fetch_one(cli, item, sem) for item in items]
        done = 0
        for coro in asyncio.as_completed(tasks):
            r = await coro
            if r: results.append(r)
            done += 1
            if done % 100 == 0 or done == total:
                e = time.time()-t0
                print(f"  {done}/{total} ({done*100//total}%) - {done/e:.1f}/s - {len(results)} ok - {e:.0f}s")
        e = time.time()-t0
        print(f"\nDone! {len(results)}/{total} in {e:.0f}s ({total/e:.1f}/s)")
    
    # Save intermediate JSON first
    json_path = OUTPUT_XLSX.parent / "nepcon_japan_2026_exhibitors.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"Saved {len(results)} records as JSON backup: {json_path}")
    
    # Save Excel
    save_xlsx(results, OUTPUT_XLSX)


if __name__ == "__main__":
    asyncio.run(main())
