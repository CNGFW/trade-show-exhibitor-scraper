#!/usr/bin/env python3
"""
electronica 2024 - FINAL scraper
Phase 1: Playwright extracts all links from list pages (with innerText parsing)
Phase 2: httpx async batch-scrapes all detail pages for phone/email/website/description
Phase 3: Classify by region and save Excel
"""
import json, re, asyncio, time
from pathlib import Path
import httpx
from playwright.async_api import async_playwright

BROWSERS_PATH = "D:/WorkBuddy/2026-06-01-10-01-40/playwright_browsers"
BASE_URL = "https://exhibitors.electronica.de/exhibitor-portal/2024/list-of-exhibitors/"
OUTPUT_DIR = Path(__file__).parent
LINKS_FILE = OUTPUT_DIR / "electronica_links_v2.json"
FULL_JSON = OUTPUT_DIR / "electronica_2024_complete.json"
OUTPUT_XLSX = OUTPUT_DIR / "electronica_2024_complete_by_region.xlsx"
CONCURRENCY = 50

CHINA_PATTERN = re.compile(r'(?:中国|china|\+86)', re.IGNORECASE)
REGION_PATTERNS = {
    "shenzhen": re.compile(r'\b(?:shenzhen|深圳|755)\b', re.IGNORECASE),
    "xiamen": re.compile(r'\b(?:xiamen|厦门|592)\b', re.IGNORECASE),
    "guangdong": re.compile(r'\b(?:dongguan|huizhou|guangzhou|zhongshan|foshan|shunde|zhuhai|东莞|惠州|广州|中山|佛山|顺德|珠海|769|752|020|760|757|756)\b', re.IGNORECASE),
    "fujian": re.compile(r'\b(?:fuzhou|xiamen|putian|sanming|quanzhou|zhangzhou|nanping|longyan|ningde|福州|厦门|莆田|三明|泉州|漳州|南平|龙岩|宁德|591|592|594|598|595|596|599|597|593)\b', re.IGNORECASE),
}
REGION_NAMES = {"shenzhen":"深圳 Shenzhen","xiamen":"厦门 Xiamen","guangdong":"广东 Guangdong","fujian":"福建 Fujian","china":"中国其他 China Other","other":"海外 Overseas"}
EXPORT_FIELDS = ["company_name","phone","email","website","url","address","city","country","postcode","booth","description","contact_person","contact_person_phone","facebook","linkedin"]
EXPORT_HEADERS = ["Company Name","Phone","Email","Website","Detail Page URL","Address","City","Country","Postcode","Booth","Description","Contact Person","Contact Person Phone","Facebook","LinkedIn"]


async def extract_links():
    """Phase 1: Extract links + basic info from list pages using innerText"""
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        page = await browser.new_page()

        print("Opening list page...")
        await page.goto(BASE_URL, wait_until="domcontentloaded", timeout=60000)
        await page.wait_for_timeout(8000)

        print("Setting 60 per page...")
        await page.evaluate("""document.querySelector('.pg_rppOpt[data-option="60"]').click()""")
        await page.wait_for_timeout(8000)

        # Go to page 1
        await page.evaluate("""() => {
            const inp = document.querySelector('[aria-label="Enter page number"]');
            inp.value = "1";
            inp.form.submit();
        }""")
        await page.wait_for_timeout(8000)

        all_items = []
        seen = set()
        page_num = 1

        while page_num <= 120:
            items = await page.evaluate("""() => {
                const divs = document.querySelectorAll('div[id^="ceId_807_"]');
                const results = [];
                divs.forEach(d => {
                    const lines = d.innerText.split('\\n').filter(l => l.trim());
                    const item = {};
                    if (lines.length >= 1) item.name = lines[0].trim();
                    if (lines.length >= 2) item.address = lines[1].trim();
                    if (lines.length >= 3) item.desc = lines[2].trim();
                    if (lines.length >= 4) item.booth = lines[3].trim();

                    // Get the detail link from within the div
                    const link = d.querySelector('a[href*="exhibitordetails"]');
                    item.url = link ? link.href : '';

                    results.push(item);
                });
                return results;
            }""")

            new_count = 0
            for item in items:
                name = item.get("name", "")
                if name and name not in seen:
                    seen.add(name)
                    # Parse address
                    addr = item.get("address", "")
                    country = ""
                    city = ""
                    postcode = ""
                    if addr:
                        parts = addr.split(",", 1)
                        if len(parts) > 1:
                            country = parts[-1].strip()
                            addr_part = parts[0].strip()
                            pc_match = re.search(r'(\d{3,6})\s+(.+)', addr_part)
                            if pc_match:
                                postcode = pc_match.group(1)
                                city = pc_match.group(2).strip()
                            else:
                                city = addr_part
                        else:
                            city = addr

                    all_items.append({
                        "company_name": name,
                        "url": item.get("url", ""),
                        "address": addr,
                        "city": city,
                        "country": country,
                        "postcode": postcode,
                        "booth": item.get("booth", ""),
                        "description": item.get("desc", ""),
                    })
                    new_count += 1

            has_next = await page.evaluate("""() => {
                for (const b of document.querySelectorAll('[name="SRField_next"]'))
                    if (!b.disabled) return true;
                return false;
            }""")

            print(f"  Page {page_num}: {len(items)} items, {new_count} new, {len(all_items)} total")

            if not has_next:
                print("  No more pages!")
                break

            await page.evaluate("""() => {
                for (const b of document.querySelectorAll('[name="SRField_next"]'))
                    if (!b.disabled) { b.click(); break; }
            }""")
            await page.wait_for_timeout(4000)
            page_num += 1

        await browser.close()
        print(f"\nTotal: {len(all_items)} exhibitors from {page_num} pages")

        with open(LINKS_FILE, "w", encoding="utf-8") as f:
            json.dump(all_items, f, ensure_ascii=False, indent=2)
        return all_items


def parse_detail_html(html, base_item):
    """Extract phone/email/website/description from detail page HTML"""
    result = dict(base_item)  # Copy basic info

    # Extract from dedicated CSS classes - look inside for <a> tags with data
    # Phone: <div class="ce_phone">...<a href="tel:...">+86 755 ...</a></div>
    phone_match = re.search(r'class="ce_phone[^"]*">.*?<a[^>]*>([^<]+)</a>', html, re.DOTALL)
    if phone_match:
        result["phone"] = phone_match.group(1).strip()

    # Mobile
    mobile_match = re.search(r'class="ce_mobile[^"]*">.*?<a[^>]*>([^<]+)</a>', html, re.DOTALL)
    if mobile_match:
        mobile = mobile_match.group(1).strip()
        if not result.get("phone"):
            result["phone"] = mobile
        else:
            result["phone"] = result["phone"] + " / " + mobile

    # Email: <div class="ce_email">...<a href="mailto:...">email@...</a></div>
    email_match = re.search(r'class="ce_email[^"]*">.*?<a[^>]*>([^<]+)</a>', html, re.DOTALL)
    if email_match:
        email = email_match.group(1).strip()
        if email not in ["info@jl-medien.de"]:
            result["email"] = email

    # Website: <div class="ce_website">...<a href="http://...">www...</a></div>
    web_match = re.search(r'class="ce_website[^"]*">.*?<a[^>]*href="([^"]+)"', html, re.DOTALL)
    if web_match:
        web = web_match.group(1).strip()
        if not any(x in web for x in ["googletagmanager", "google-analytics"]):
            result["website"] = web.rstrip('/')

    # Address: <div class="ce_addr">...</div>
    addr_match = re.search(r'class="ce_addr[^"]*">\s*(.*?)\s*</div>', html, re.DOTALL)
    if addr_match:
        addr = re.sub(r'<[^>]+>', '', addr_match.group(1)).strip()
        if addr:
            result["address"] = addr

    # Description from ce_text (first one, before nested elements)
    desc_match = re.search(r'class="ce_text[^"]*">(.*?)</div>\s*<(?:div|/div)', html, re.DOTALL)
    if desc_match:
        desc = re.sub(r'<[^>]+>', '', desc_match.group(1)).strip()
        if len(desc) > 10:
            result["description"] = desc

    # Social links
    result["facebook"] = ""
    result["linkedin"] = ""
    for m in re.finditer(r'href="(https?://(?:www\.)?(?:facebook|linkedin)\.com[^"]*)"', html, re.I):
        url = m.group(1)
        if 'facebook.com' in url: result["facebook"] = url
        if 'linkedin.com' in url: result["linkedin"] = url

    # Contact person from second ce_cntct section
    result["contact_person"] = ""
    result["contact_person_phone"] = ""
    # Find sections with ce_topic (indicates a contact person, not the main company)
    contact_sections = re.finditer(
        r'class="ce_cntct[^"]*">.*?</div>\s*</div>\s*</div>',
        html, re.DOTALL
    )
    for sec_match in contact_sections:
        sec = sec_match.group()
        # Check if it has a ce_topic (contact person section)
        topic_match = re.search(r'class="ce_topic[^"]*">\s*(.*?)\s*</div>', sec, re.DOTALL)
        if topic_match:
            topic = re.sub(r'<[^>]+>', '', topic_match.group(1)).strip()
            # Name from ce_head
            head_match = re.search(r'class="ce_head[^"]*">\s*(.*?)\s*</div>', sec, re.DOTALL)
            if head_match:
                name = re.sub(r'<[^>]+>', '', head_match.group(1)).strip()
                result["contact_person"] = f"{topic}: {name}"
            # Position from ce_pstn
            pstn_match = re.search(r'class="ce_pstn[^"]*">\s*(.*?)\s*</div>', sec, re.DOTALL)
            if pstn_match:
                pstn = re.sub(r'<[^>]+>', '', pstn_match.group(1)).strip()
                result["contact_person"] += f" ({pstn})"
            # Phone
            ph_match = re.search(r'class="ce_phone[^"]*">.*?<a[^>]*>([^<]+)</a>', sec, re.DOTALL)
            if ph_match:
                result["contact_person_phone"] = ph_match.group(1).strip()
            break  # Only first contact person

    # Better description from detail page
    desc_match = re.search(r'class="ce_text">(.*?)</div>', html, re.DOTALL)
    if desc_match:
        desc = re.sub(r'<[^>]+>', '', desc_match.group(1)).strip()
        if len(desc) > len(result.get("description", "")):
            result["description"] = desc

    return result


async def scrape_details(items):
    """Phase 2: Batch scrape detail pages via httpx"""
    total = len(items)
    results = []
    sem = asyncio.Semaphore(CONCURRENCY)

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
        "Accept": "text/html,application/xhtml+xml",
    }

    async with httpx.AsyncClient(headers=headers, follow_redirects=True, timeout=20) as client:
        print(f"\nScraping {total} detail pages ({CONCURRENCY} concurrent)...")
        t0 = time.time()

        async def fetch_one(i, item):
            url = item.get("url", "")
            if not url:
                results.append(item)
                return
            for _ in range(2):
                try:
                    async with sem:
                        r = await client.get(url)
                    if r.status_code == 200:
                        results.append(parse_detail_html(r.text, item))
                        return
                except:
                    await asyncio.sleep(1)
            results.append(item)  # fallback to basic info

        tasks = [fetch_one(i, item) for i, item in enumerate(items)]
        done = 0
        for coro in asyncio.as_completed(tasks):
            await coro
            done += 1
            if done % 200 == 0 or done == total:
                elapsed = time.time() - t0
                print(f"  {done}/{total} ({done*100//total}%) - {elapsed:.0f}s")

        print(f"Done! {total} pages in {time.time()-t0:.0f}s")
        return results


def classify(item):
    c = ' '.join(str(item.get(k,'') or '') for k in ['company_name','address','city','country','phone','email','website','description'])
    if not CHINA_PATTERN.search(c): return "other"
    if REGION_PATTERNS["shenzhen"].search(c): return "shenzhen"
    if REGION_PATTERNS["xiamen"].search(c): return "xiamen"
    if REGION_PATTERNS["guangdong"].search(c): return "guangdong"
    if REGION_PATTERNS["fujian"].search(c): return "fujian"
    return "china"


def sanitize(v):
    if not isinstance(v, str): return v if v is not None else ""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', v)


def save_xlsx(classified, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook(); wb.remove(wb.active)
    order = ["shenzhen","xiamen","guangdong","fujian","china","other"]

    for rk in order:
        items = classified.get(rk, [])
        if not items: continue
        ws = wb.create_sheet(title=REGION_NAMES.get(rk, rk)[:31])
        hf = Font(bold=True, color="FFFFFF", size=10)
        hfill = PatternFill(start_color="005BAC", end_color="005BAC", fill_type="solid")
        tb = Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
        for c, h in enumerate(EXPORT_HEADERS, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = hf; cell.fill = hfill; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = tb
        for ri, item in enumerate(items, 2):
            for ci, f in enumerate(EXPORT_FIELDS, 1):
                val = sanitize(item.get(f, "") or "")
                ws.cell(row=ri, column=ci, value=val).font = Font(size=9)
                ws.cell(row=ri, column=ci).alignment = Alignment(vertical="top", wrap_text=True)
                ws.cell(row=ri, column=ci).border = tb
        widths = [35,20,30,30,60,40,20,20,12,15,50,30,30,20]
        for i, w in enumerate(widths): ws.column_dimensions[chr(65+i)].width = w
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet(title="Summary", index=0)
    ws["A1"] = "electronica 2024 - Complete Exhibitor Data"
    ws["A1"].font = Font(bold=True, size=14)
    total = sum(len(classified.get(r, [])) for r in order)
    ws["A3"] = "Total Exhibitors:"; ws["B3"] = total
    for i, rk in enumerate(order, 5):
        ws[f"A{i}"] = REGION_NAMES.get(rk, rk); ws[f"B{i}"] = len(classified.get(rk, []))
    ws.column_dimensions["A"].width = 35; ws.column_dimensions["B"].width = 12

    wb.save(path)
    print(f"Saved to {path}")


async def main():
    # Phase 1
    if LINKS_FILE.exists():
        with open(LINKS_FILE, "r", encoding="utf-8") as f:
            items = json.load(f)
        print(f"Loaded {len(items)} links from {LINKS_FILE}")
    else:
        items = await extract_links()

    # Phase 2
    results = await scrape_details(items)

    # Save full data
    with open(FULL_JSON, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"Full data saved to {FULL_JSON}")

    # Phase 3: Classify
    classified = {}
    for item in results:
        rk = classify(item)
        classified.setdefault(rk, []).append(item)

    print("\n=== Classification ===")
    for rk in ["shenzhen","xiamen","guangdong","fujian","china","other"]:
        n = len(classified.get(rk, []))
        print(f"  {REGION_NAMES.get(rk,rk):35s}: {n:4d}")

    save_xlsx(classified, OUTPUT_XLSX)


if __name__ == "__main__":
    asyncio.run(main())
