"""
EUROGUSS 2026 Exhibitor Scraper
Data source: Algolia Search API (prod_website_companies_en index, filter: site:guss)
Total: 728 exhibitors
"""
import json, re, time, requests, logging
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

OUTPUT_DIR = Path(__file__).parent

# ── Algolia API Config ──
API_URL = "https://4eb6g0v1nt-2.algolianet.com/1/indexes/*/queries"
HEADERS = {
    "accept": "application/json",
    "content-type": "text/plain",
    "x-algolia-api-key": "f0416e3d1b38ae3aa789c8750e12bfe5",
    "x-algolia-application-id": "4EB6G0V1NT",
}
INDEX_NAME = "prod_website_companies_en"
BASE_URL = "https://www.euroguss.de"

# ── Region Classification ──
def classify_region(country, city, company_name):
    """Classify Chinese exhibitors by region"""
    if country not in ("China", "Hong Kong SAR China", "Taiwan", "Taiwan, China"):
        return "Overseas"
    city_lower = (city or "").lower()
    name_lower = (company_name or "").lower()
    # Shenzhen
    if "shenzhen" in city_lower or "shenzhen" in name_lower or "深圳" in name_lower:
        return "Shenzhen"
    # Xiamen
    if "xiamen" in city_lower or "xiamen" in name_lower or "厦门" in name_lower:
        return "Xiamen"
    # Guangdong province (but not Shenzhen)
    if "guangdong" in city_lower or "guangzhou" in city_lower or "dongguan" in city_lower or \
       "foshan" in city_lower or "zhuhai" in city_lower or "huizhou" in city_lower or \
       "zhongshan" in city_lower or "jiangmen" in city_lower:
        return "Guangdong"
    # Fujian (but not Xiamen)
    if "fujian" in city_lower or "fuzhou" in city_lower or "quanzhou" in city_lower or \
       "zhangzhou" in city_lower:
        return "Fujian"
    # Zhejiang
    if "zhejiang" in city_lower or "hangzhou" in city_lower or "ningbo" in city_lower or \
       "wenzhou" in city_lower or "yiwu" in city_lower:
        return "Zhejiang"
    # Jiangsu
    if "jiangsu" in city_lower or "suzhou" in city_lower or "nanjing" in city_lower or \
       "wuxi" in city_lower or "changzhou" in city_lower or "kunshan" in city_lower:
        return "Jiangsu"
    # Shanghai
    if "shanghai" in city_lower or "shanghai" in name_lower:
        return "Shanghai"
    # Hong Kong / Taiwan
    if country == "Hong Kong SAR China":
        return "Hong Kong"
    if country == "Taiwan" or country == "Taiwan, China":
        return "Taiwan"
    # Other China
    return "China Other"


def clean_text(text):
    """Strip HTML tags, entities, control chars, collapse whitespace"""
    if not text:
        return ""
    text = str(text)
    text = re.sub(r'<[^>]+>', '', text)
    text = re.sub(r'&[a-z0-9]+;', '', text)
    text = re.sub(r'\s+', ' ', text).strip()
    text = re.sub(r'[\x00-\x1f\x7f-\x9f]', '', text)
    return text


def fetch_all_exhibitors():
    """Fetch all 728 exhibitors from Algolia in one request"""
    payload = {
        "requests": [{
            "indexName": INDEX_NAME,
            "params": (
                "distinct=true"
                "&facets=%5B%22*%22%5D"
                "&filters=site%3Aguss"
                "&hitsPerPage=1000"
                "&page=0"
                "&query="
            )
        }]
    }
    
    logger.info("Fetching exhibitors from Algolia API...")
    resp = requests.post(API_URL, json=payload, headers=HEADERS, timeout=30)
    resp.raise_for_status()
    data = resp.json()
    
    results = data.get("results", [{}])[0]
    hits = results.get("hits", [])
    total = results.get("nbHits", 0)
    
    logger.info(f"Got {len(hits)} hits, API reports {total} total")
    return hits


def extract_exhibitor(hit):
    """Extract structured data from an Algolia hit"""
    booth_info = (hit.get("booth") or [{}])[0]
    
    # Employee list
    employees = hit.get("employee") or []
    employee_names = ", ".join(
        f"{e.get('firstName', '')} {e.get('lastName', '')}".strip()
        for e in employees
    )
    employee_functions = ", ".join(
        e.get("function", "") for e in employees if e.get("function")
    )
    employee_emails = ", ".join(
        e.get("email", "") for e in employees if e.get("email")
    )
    
    # Categories
    def_cats = [v for k in hit.get("filternomenclature_DEF", {}) for v in (hit["filternomenclature_DEF"].get(k) or [])]
    branch_cats = [v for k in hit.get("filternomenclature_BRANCHE", {}) for v in (hit["filternomenclature_BRANCHE"].get(k) or [])]
    beruf_cats = [v for k in hit.get("filternomenclature_BERUF", {}) for v in (hit["filternomenclature_BERUF"].get(k) or [])]
    
    # Country
    country = hit.get("country", "")
    
    # Detail page URL
    detail_url = f"{BASE_URL}{hit.get('url', '')}"
    
    # Region classification
    region = classify_region(country, hit.get("city", ""), hit.get("companyName", ""))
    
    return {
        "Company Name": clean_text(hit.get("companyName", "")),
        "Country": country,
        "Region": region,
        "Hall": booth_info.get("boothHall", ""),
        "Booth Number": booth_info.get("boothNumber", ""),
        "Company Type": hit.get("companyType") or "",
        "Street": hit.get("streetno", ""),
        "Postcode": hit.get("postcode", ""),
        "City": hit.get("city", ""),
        "Email": hit.get("email", ""),
        "Website": "",  # Not in Algolia data
        "Phone": "",    # Not in Algolia data
        "Description": clean_text(hit.get("companyDescription") or ""),
        "Products": clean_text(", ".join(hit.get("products") or [])),
        "Keywords": clean_text(", ".join(hit.get("keyword") or [])),
        "Categories (DEF)": ", ".join(def_cats),
        "Categories (BRANCHE)": ", ".join(branch_cats),
        "Categories (BERUF)": ", ".join(beruf_cats),
        "Co-Exhibitors": ", ".join(hit.get("coExhibitors") or []),
        "Contact Names": employee_names,
        "Contact Functions": employee_functions,
        "Contact Emails": employee_emails,
        "Logo URL": hit.get("logo") or "",
        "Slogan": hit.get("slogan") or "",
        "Detail URL": detail_url,
        "Object ID": hit.get("objectID", ""),
    }


def save_excel(exhibitors, filename):
    """Save to Excel with formatting"""
    wb = Workbook()
    ws = wb.active
    ws.title = "EUROGUSS Exhibitors"
    
    if not exhibitors:
        logger.warning("No exhibitors to save!")
        return
    
    headers = list(exhibitors[0].keys())
    
    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    
    # Data rows
    for row_idx, ex in enumerate(exhibitors, 2):
        for col_idx, header in enumerate(headers, 1):
            val = ex.get(header, "")
            # Sanitize: remove control characters
            if isinstance(val, str):
                val = re.sub(r'[\x00-\x1f\x7f-\x9f]', '', val)
            ws.cell(row=row_idx, column=col_idx, value=val)
    
    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col[:50]:  # Check first 50 rows
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)
    
    # Freeze header
    ws.freeze_panes = "A2"
    
    # Auto-filter
    ws.auto_filter.ref = f"A1:{col_letter}{len(exhibitors) + 1}"
    
    # ── Region Summary Sheet ──
    ws2 = wb.create_sheet("Region Summary")
    from collections import Counter
    region_counts = Counter(ex.get("Region", "Unknown") for ex in exhibitors)
    
    ws2.append(["Region", "Count"])
    for region, count in region_counts.most_common():
        ws2.append([region, count])
    ws2.append(["TOTAL", sum(region_counts.values())])
    
    # Summary sheet styling
    for cell in ws2["A1:B1"][0]:
        cell.font = header_font
        cell.fill = header_fill
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 12
    
    filepath = OUTPUT_DIR / filename
    wb.save(filepath)
    logger.info(f"Saved Excel: {filepath}")
    return filepath


def main():
    # 1. Fetch
    hits = fetch_all_exhibitors()
    
    # 2. Extract
    exhibitors = [extract_exhibitor(h) for h in hits]
    logger.info(f"Extracted {len(exhibitors)} exhibitors")
    
    # 3. Save JSON backup
    json_path = OUTPUT_DIR / "euroguss_exhibitors.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(exhibitors, f, ensure_ascii=False, indent=2)
    logger.info(f"Saved JSON: {json_path}")
    
    # 4. Save Excel
    xlsx_path = save_excel(exhibitors, "euroguss_exhibitors.xlsx")
    
    # 5. Print region summary
    from collections import Counter
    region_counts = Counter(ex.get("Region", "Unknown") for ex in exhibitors)
    logger.info("Region Distribution:")
    for region, count in region_counts.most_common():
        logger.info(f"  {region}: {count}")
    
    # 6. Print country summary
    country_counts = Counter(ex.get("Country", "Unknown") for ex in exhibitors)
    logger.info("Top 10 Countries:")
    for country, count in country_counts.most_common(10):
        logger.info(f"  {country}: {count}")
    
    return xlsx_path


if __name__ == "__main__":
    main()
